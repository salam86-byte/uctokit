"""Testy extrakce z tabulkové faktury (XLSX)."""

import datetime
import io
import unittest
from datetime import date
from decimal import Decimal

from uctokit.invoices import pipeline as P
from uctokit.invoices import xlsx_text
from uctokit.invoices.types import SourceDocument


def _xlsx(rows) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class XlsxTextTests(unittest.TestCase):
    def test_flatten_cells_to_text(self):
        data = _xlsx([
            ["Faktura č.", "2026-77"],
            ["IČO:", "12345679"],
            ["Celkem k úhradě", "12 100,00 Kč"],
        ])
        text = xlsx_text.extract_text(data)
        for token in ("2026-77", "12345679", "12 100,00"):
            self.assertIn(token, text)

    def test_empty_on_garbage(self):
        self.assertEqual(xlsx_text.extract_text(b"nonsense"), "")

    def test_preserves_displayed_decimal_amount(self):
        from openpyxl import Workbook

        wb = Workbook()
        cell = wb.active["A1"]
        cell.value = 4500
        cell.number_format = '#,##0.00 "Kč"'
        buf = io.BytesIO()
        wb.save(buf)
        self.assertIn("4500.00", xlsx_text.extract_text(buf.getvalue()))


class XlsxPipelineTests(unittest.TestCase):
    def test_routes_through_xlsx_and_heuristics(self):
        data = _xlsx([
            ["Dodavatel: Test s.r.o., IČO: 12345679"],
            ["Variabilní symbol 202677"],
            ["Celkem k úhradě: 12 100,00 Kč"],
        ])
        doc = SourceDocument(content=data, filename="faktura.xlsx")
        res = P.extract(doc)  # bez LLM – jen heuristiky nad zploštěným textem
        self.assertEqual(res.method, "xlsx")
        self.assertEqual(res.invoice.supplier_ico.value, "12345679")
        self.assertEqual(res.invoice.variable_symbol.value, "202677")

    def test_regression_separate_bank_code_and_integer_amount(self):
        """Rozložení reálné trenérské faktury: dodavatel vlevo, klub vpravo."""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        # Smyšlené údaje se stejnou strukturou jako reálná trenérská faktura.
        ws["B1"], ws["D1"] = "Jan Vzorník", "FAKTURA"
        ws["D2"], ws["E2"] = "DATUM:", "17.3.2026"
        ws["D3"], ws["E3"] = "Č. FAKTURY", 16
        ws["B9"], ws["E9"] = "IČ:12345670", "IČO: 87654321"
        ws["B14"], ws["E14"] = "Trénink v měsíci Březen 03/2026", 4500
        ws["E14"].number_format = '#,##0.00 "Kč"'
        ws["B29"] = "SPLATNOST FAKTURY 27.3.2026"
        ws["B32"] = "Číslo účtu                 1234567890"
        ws["B33"] = "Identifikace banky         0800"
        buf = io.BytesIO()
        wb.save(buf)

        res = P.extract(SourceDocument(buf.getvalue(), "faktura.xlsx"))
        inv = res.invoice
        self.assertEqual(inv.supplier_name.value, "Jan Vzorník")
        self.assertEqual(inv.supplier_ico.value, "12345670")
        self.assertEqual(inv.supplier_account.value, "1234567890/0800")
        self.assertEqual(inv.total_amount.value, Decimal("4500.00"))
        self.assertEqual(inv.invoice_number.value, "16")
        self.assertEqual(inv.issue_date.value, date(2026, 3, 17))
        self.assertEqual(inv.due_date.value, date(2026, 3, 27))



class XlsTests(unittest.TestCase):
    """Starý binární .xls (Excel 97–2003) — čte se xlrd, ne openpyxl."""

    def _workbook(self):
        """Sešit s textem, číslem a DATEM uloženým jako pořadové číslo."""
        xlwt = __import__("xlwt")
        wb = xlwt.Workbook()
        ws = wb.add_sheet("List1")
        date_style = xlwt.XFStyle()
        date_style.num_format_str = "DD.MM.YYYY"
        ws.write(0, 0, "FAKTURA")
        ws.write(1, 0, "Faktura č.:")
        ws.write(1, 1, 25003)
        ws.write(2, 0, "Den splatnosti:")
        ws.write(2, 1, datetime.date(2026, 1, 17), date_style)
        ws.write(3, 0, "K úhradě celkem:")
        ws.write(3, 1, 650.0)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_reads_text_numbers_and_dates(self):
        try:
            import xlwt  # noqa: F401
        except ImportError:
            self.skipTest("xlwt není k dispozici (jen pro sestavení testovacího .xls)")
        text = xlsx_text.extract_text_xls(self._workbook())
        self.assertIn("FAKTURA", text)
        self.assertIn("25003", text)
        self.assertIn("650", text)
        # Datum MUSÍ být datum, ne pořadové číslo 46039 — jinak ho extrakce nenajde.
        self.assertIn("2026-01-17", text)
        self.assertNotIn("46039", text)

    def test_broken_file_returns_empty(self):
        self.assertEqual(xlsx_text.extract_text_xls(b"tohle neni xls"), "")

if __name__ == "__main__":
    unittest.main()
